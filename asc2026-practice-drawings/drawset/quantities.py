"""
Quantity answer key for the practice set (computed from drawset/model.py).

    python3 -m drawset.quantities          -> output/<name>_Quantity_Answer_Key.xlsx

Quantities are "neat" (no waste) unless a waste column says otherwise. The intent is that
students do their own takeoff first and then compare. Formulas are left live in Excel.
"""
from __future__ import annotations

import math
import os
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from shapely.geometry import box
from shapely.ops import unary_union

from . import model as M
from . import plans as PL

IN = M.IN
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(HERE, "output")

HDR = PatternFill("solid", fgColor="D9D9D9")
SUB = PatternFill("solid", fgColor="F2F2F2")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class Sheet:
    def __init__(self, wb, title, cols, widths):
        self.ws = wb.create_sheet(title)
        self.row = 1
        self.cols = cols
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.title(title.upper())
        self.header(cols)

    def title(self, text):
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(bold=True, size=13)
        self.row += 2

    def header(self, cols):
        for i, t in enumerate(cols, 1):
            c = self.ws.cell(self.row, i, t)
            c.font = BOLD
            c.fill = HDR
            c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        self.ws.row_dimensions[self.row].height = 30
        self.row += 1

    def section(self, text):
        c = self.ws.cell(self.row, 1, text)
        c.font = BOLD
        for i in range(1, len(self.cols) + 1):
            self.ws.cell(self.row, i).fill = SUB
        self.row += 1

    def add(self, values, fmt=None, bold=False):
        r = self.row
        for i, v in enumerate(values, 1):
            if isinstance(v, str) and "{r}" in v:
                v = v.replace("{r}", str(r))
            c = self.ws.cell(r, i, v)
            c.border = BOX
            if bold:
                c.font = BOLD
            if fmt and i - 1 < len(fmt) and fmt[i - 1]:
                c.number_format = fmt[i - 1]
        self.row += 1
        return r

    def blank(self, n=1):
        self.row += n


# ----------------------------------------------------------------------------------------
# geometry helpers
# ----------------------------------------------------------------------------------------
OUT_FACE = M.EW_OUT
BRICK_BOT = 99.0
BRICK_TOP = 131.0                 # underside of coping (coping covers the top 4")
LINK_BRICK_TOP = 116 + 4 * IN
CMU_BOT = M.LEVELS["GRADE"]       # top of foundation wall 99'-4"
CMU_TOP = M.LEVELS["PARAPET"]
LINK_CMU_TOP = M.LEVELS["LINK_PARAPET"]


def facade_segments():
    """(name, outside length, bottom, top, notes) for every brick face."""
    L_long = 150 + 2 * OUT_FACE
    L_short = 72 + 2 * OUT_FACE
    link_w = 12 + 2 * OUT_FACE
    segs = [
        ("NORTH (A-201/1)", L_long, BRICK_BOT, BRICK_TOP, ""),
        ("SOUTH (A-201/2)", L_long, BRICK_BOT, BRICK_TOP, ""),
        ("EAST (A-202/1)", L_short, BRICK_BOT, BRICK_TOP, ""),
        ("WEST (A-202/2) GROSS", L_short, BRICK_BOT, BRICK_TOP, ""),
        ("WEST - LESS AREA BEHIND LINK", -link_w, BRICK_BOT, LINK_BRICK_TOP, "link abuts main wall"),
        ("LINK NORTH (A-202/3)", 36.0 - OUT_FACE, BRICK_BOT, LINK_BRICK_TOP,
         "existing face x -36 to main brick face"),
        ("LINK SOUTH (A-202/4)", 36.0 - OUT_FACE, BRICK_BOT, LINK_BRICK_TOP,
         "existing face x -36 to main brick face"),
    ]
    return segs


def openings_on_face(face):
    m = {"NORTH": ("N",), "SOUTH": ("S",), "EAST": ("E",), "WEST": ("W1", "W2", "W0"),
         "LINK NORTH": ("LN",), "LINK SOUTH": ("LS",)}[face]
    out = []
    for o in M.OPENINGS:
        if o.wall in m and o.kind in ("window", "storefront", "louver", "door"):
            if o.kind == "door" and o.frame == "SF":
                continue  # inside SF-1
            out.append(o)
    return out


def interior_cmu(level, wtype):
    """net plan area of interior CMU of a type (excluding exterior CMU), and effective LF"""
    ext = M.ext_layer(level, -M.EW_IN, M.EW_OUT)
    parts = [w.poly() for w in M.interior_walls(level, types=(wtype,))]
    others = [w.poly() for w in M.interior_walls(level) if w.type != wtype and
              M.WALL_TYPES[w.type]["mat"] == "cmu"]
    g = unary_union(parts).difference(ext)
    if wtype == "P1" and others:
        # junctions shared with P2 are counted as P2
        g = g.difference(unary_union(others))
    return g.area, g.area / M.CMU_T


def wall_height(level):
    return (113 + 4 * IN - 100.0) if level == "L1" else (127 + 4 * IN - 114.0)


# ----------------------------------------------------------------------------------------
def build():
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "README"
    lines = [
        "CEDAR PRAIRIE ES ADDITION - ESTIMATING PRACTICE SET - QUANTITY ANSWER KEY",
        "",
        "Fictional project. Quantities are computed directly from the drawing model (drawset/model.py).",
        "They are NEAT quantities (no waste) unless a waste column is shown. Use them only to check your own takeoff.",
        "Typical takeoff conventions used:",
        " - Masonry areas: outside face length x height, less masonry openings (MO). Brick from 99'-0\" to underside of coping 131'-0\" (link 116'-4\").",
        " - Exterior CMU: grid (centerline) length x 99'-4\" to 131'-4\" (link 99'-4\" to 116'-8\"), less MO.",
        " - Interior CMU: net plan area / 7 5/8\" = effective centerline length (junctions counted once) x height to structure, less door MO.",
        "   L1 height 13'-4\" (100'-0\" to 113'-4\"), L2 height 13'-4\" (114'-0\" to 127'-4\").",
        " - Flooring: clear room area inside finished wall faces, by pattern zone shown on A-111 / A-112 (no deduction for casework).",
        " - Base: room perimeter less door openings. Wall tile: perimeter x 7'-0\" less door openings.",
        " - Doors/frames/hardware and windows/storefront: counts straight from the schedules.",
        "Differences of 1-3% from a careful manual takeoff are normal (corner/junction conventions).",
    ]
    for i, t in enumerate(lines, 1):
        ws0.cell(i, 1, t).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    ws0.column_dimensions["A"].width = 140

    # ===================================== MASONRY =====================================
    S = Sheet(wb, "Masonry", ["ITEM", "DESCRIPTION", "LENGTH (FT)", "HEIGHT (FT)", "QTY", "AREA (SF)",
                              "UNIT", "NOTES"], [34, 50, 12, 12, 8, 14, 8, 40])
    fm = [None, None, "0.00", "0.00", "0", "#,##0.0", None, None]
    S.section("A. FACE BRICK (FB-1 FIELD + FB-2 ACCENT) - GROSS FACADE AREA")
    first = S.row
    for name, L, b, t, note in facade_segments():
        S.add([name, "GROSS WALL FACE", L, t - b, 1, "=C{r}*D{r}*E{r}", "SF", note], fm)
    last = S.row - 1
    gross_row = S.add(["GROSS BRICK FACE", "", "", "", "", f"=SUM(F{first}:F{last})", "SF", ""], fm, bold=True)
    S.section("B. DEDUCT MASONRY OPENINGS (MO)")
    first = S.row
    for face in ("NORTH", "SOUTH", "EAST", "WEST", "LINK NORTH", "LINK SOUTH"):
        cnt = Counter()
        for o in openings_on_face(face):
            h = o.head - o.sill
            if o.wall == "W0" or (o.type == "W-C"):
                pass
            cnt[(o.type if o.kind != "door" else f"DOOR {o.frame}", round(o.w, 4), round(h, 4))] += 1
        for (typ, w, h), n in sorted(cnt.items()):
            S.add([f"{face} - {typ}", "OPENING (MO)", w, h, n, "=-C{r}*D{r}*E{r}", "SF", ""], fm)
    last = S.row - 1
    open_row = S.add(["TOTAL OPENINGS", "", "", "", "", f"=SUM(F{first}:F{last})", "SF", ""], fm, bold=True)
    S.section("C. CAST STONE (DEDUCT FROM BRICK)")
    perim_out = 2 * (150 + 2 * OUT_FACE) + 2 * (72 + 2 * OUT_FACE)
    link_w = 12 + 2 * OUT_FACE
    cs2_lf = perim_out - 2 * 4.0 - link_w
    r_cs2 = S.add(["CS-2 CAST STONE BAND @ L2", "8\" HIGH, 113'-4\" TO 114'-0\"", cs2_lf, 8 * IN, 1,
                   "=-C{r}*D{r}*E{r}", "SF", "perimeter less 2 stair W-C strips less link"], fm)
    sills = [o for o in M.OPENINGS if o.type in ("W-A", "W-B", "SF-2")]
    cs1 = Counter(round(o.w, 4) for o in sills)
    first = S.row
    for w, n in sorted(cs1.items()):
        S.add([f"CS-1 SILL {w:.2f}' MO", "4\" HIGH x MO WIDTH + 4\" (2\" EACH END)", w + 4 * IN, 4 * IN, n,
               "=-C{r}*D{r}*E{r}", "SF", "count = windows of this width"], fm)
    last = S.row - 1
    S.add(["NET FACE BRICK", "GROSS - OPENINGS - CAST STONE", "", "", "",
           f"=F{gross_row}+F{open_row}+F{r_cs2}+SUM(F{first}:F{last})", "SF", "FB-1 + FB-2"], fm, bold=True)
    net_brick_row = S.row - 1
    # FB-2 accent split
    base_lf = perim_out - link_w + 2 * 36.0
    door_w_grade = sum(o.w for o in M.OPENINGS if o.wall in M.EXT_SEGS and o.sill <= 100.01 and
                       o.kind in ("door", "storefront") and not (o.kind == "door" and o.frame == "SF"))
    S.add(["FB-2 BASE BAND", "6 COURSES 99'-4\" TO 100'-8\" (+1 COURSE BELOW GRADE)", base_lf - door_w_grade,
           20 * IN, 1, "=C{r}*D{r}*E{r}", "SF", "99'-0\" to 100'-8\"; less openings at grade"], fm)
    sold = [o for o in M.OPENINGS if o.type in ("W-A", "W-B", "SF-2")]
    sold_lf = sum(o.w + 8 * IN for o in sold)
    S.add(["FB-2 SOLDIER COURSE AT HEADS", f"{len(sold)} OPENINGS, MO + 4\" EACH SIDE", sold_lf, 8 * IN, 1,
           "=C{r}*D{r}*E{r}", "SF", "W-A, W-B, SF-2 heads"], fm)
    S.add(["MODULAR BRICK COUNT (REFERENCE)", "6.75 BRICK / SF (3/8\" JOINTS), NO WASTE", "", "", "",
           f"=F{net_brick_row}*6.75", "EA", "add 3-5% waste"], fm)
    S.add(["CS-2 BAND LENGTH", "", cs2_lf, "", "", "=C{r}", "LF", ""], fm)
    S.add(["CS-1 SILLS", "", "", "", len(sills), "", "EA", f"widths: {', '.join(f'{w:.2f}' for w in sorted(cs1))}"], fm)
    S.blank()
    # exterior CMU backup
    S.section("D. EXTERIOR 8\" CMU BACKUP (EW-1) - GRID LENGTH x HEIGHT")
    first = S.row
    S.add(["MAIN BLOCK PERIMETER", "GRID LINES 1-6 / A-D", 2 * (150 + 72), CMU_TOP - CMU_BOT, 1,
           "=C{r}*D{r}*E{r}", "SF", "99'-4\" to 131'-4\""], fm)
    S.add(["LESS: LINK OPENING AT GRID 1 (L1)", "WALL STARTS ON L2 FRAMING ABOVE LINK", -12.0, 114.0 - CMU_BOT, 1,
           "=C{r}*D{r}*E{r}", "SF", "x = 0, y 30-42"], fm)
    S.add(["LINK NORTH + SOUTH WALLS", "x -36 TO 0", 72.0, LINK_CMU_TOP - CMU_BOT, 1, "=C{r}*D{r}*E{r}", "SF", ""], fm)
    ext_open = 0.0
    for o in M.OPENINGS:
        if o.wall in M.EXT_SEGS and o.kind in ("window", "storefront", "louver", "door") and \
                not (o.kind == "door" and o.frame == "SF"):
            ext_open += o.w * (o.head - o.sill)
    S.add(["LESS: EXTERIOR OPENINGS (MO)", "ALL WINDOWS, STOREFRONT, LOUVERS, DOORS", ext_open, 1, -1,
           "=C{r}*D{r}*E{r}", "SF", "same MO list as section B"], fm)
    last = S.row - 1
    S.add(["NET EXTERIOR CMU", "", "", "", "", f"=SUM(F{first}:F{last})", "SF", "x 1.125 units/SF"], fm, bold=True)
    S.blank()
    S.section("E. INTERIOR 8\" CMU PARTITIONS (P1 / P2)")
    for level in ("L1", "L2"):
        H = wall_height(level)
        for wt in ("P1", "P2"):
            a, lf = interior_cmu(level, wt)
            first = S.add([f"{level} {wt}", M.WALL_TYPES[wt]["desc"][:48], lf, H, 1, "=C{r}*D{r}*E{r}", "SF",
                           "effective centerline LF"], fm)
            dsum = 0.0
            n = 0
            for o in M.OPENINGS:
                if o.level != level or o.wall in M.EXT_SEGS:
                    continue
                w = M.WALL_BY_ID.get(o.wall)
                if w is None or w.type != wt:
                    continue
                dsum += o.w * (o.head - o.sill)
                n += 1
            S.add([f"{level} {wt} LESS DOOR / ELEV. MO", f"{n} OPENINGS", dsum, 1, -1, "=C{r}*D{r}*E{r}", "SF", ""], fm)
            S.add([f"{level} {wt} NET", "", "", "", "", f"=F{first}+F{first + 1}", "SF", ""], fm, bold=True)
    S.add(["NOTE", "CMU units = SF x 1.125; mortar about 8.5 CF per 100 SF; joint reinforcing @ 16\" = 0.75 LF/SF", "", "", "",
           "", "", ""])
    S.blank()
    S.section("F. LINTELS (COUNT BY OPENING WIDTH) - SEE S-302 SCHEDULE")
    cnt = Counter()
    for o in M.OPENINGS:
        if o.kind == "elevator":
            continue
        if o.kind == "door" and o.frame == "SF":
            continue
        if o.wall in M.EXT_SEGS:
            cnt[("EXTERIOR (BRICK ANGLE + CMU LINTEL)", round(o.w, 3))] += 1
        else:
            w = M.WALL_BY_ID[o.wall]
            if M.WALL_TYPES[w.type]["mat"] == "cmu":
                cnt[("INTERIOR CMU LINTEL", round(o.w, 3))] += 1
    for o in M.OPENINGS:
        if o.kind == "elevator":
            cnt[("INTERIOR CMU LINTEL (ELEVATOR ENTRANCE)", round(o.w, 3))] += 1
    for (k, w), n in sorted(cnt.items()):
        S.add([k, f"MO {w:.3f} FT", w, "", n, "", "EA", ""], fm)

    # ===================================== FLOORING =====================================
    from .sheets_arch_plans import floor_zones
    F = Sheet(wb, "Flooring", ["LEVEL", "ROOM", "FINISH", "AREA (SF)", "WASTE %", "ORDER QTY (SF)", "NOTES"],
              [8, 26, 12, 12, 9, 15, 50])
    ff = [None, None, None, "#,##0.0", "0%", "#,##0", None]
    waste = {"VCT-1": 0.07, "VCT-2": 0.15, "LVT-1": 0.08, "LVT-2": 0.12, "CPT-1": 0.06, "WOM-1": 0.08,
             "PT-1": 0.10, "RF-1": 0.10, "RST-1": 0.0, "SC-1": 0.0}
    totals = defaultdict(float)
    for level in ("L1", "L2"):
        F.section(f"LEVEL {level[-1]}")
        for code, g, r in floor_zones(level):
            if g.is_empty or code == "RST-1":
                continue
            if code == "RF-1" and level == "L2":
                # the intermediate landing is already counted on level 1
                sg = PL.stair_geom(r.num)
                y0, y1 = sg["floor_land"]
                g = g.intersection(box(sg["x0"] - 1, y0 - 0.01, sg["x1"] + 1, y1 + 0.01))
            totals[code] += g.area
            F.add([level, f"{r.num} {r.name.replace(chr(10), ' ')}", code, round(g.area, 1), waste.get(code, 0.05),
                   "=D{r}*(1+E{r})", ""], ff)
    F.section("TOTALS BY FINISH (NEAT)")
    for code in sorted(totals):
        F.add(["", "TOTAL", code, round(totals[code], 1), waste.get(code, 0.05), "=D{r}*(1+E{r})",
               M.FINISHES.get(code, "")], ff, bold=True)
    F.blank()
    F.section("STAIRS (RST-1)")
    risers = 24 * 2
    treads = 22 * 2
    sg = PL.stair_geom("ST-1")
    F.add(["", "ST-1 + ST-2", "RST-1", "", "", "", f"{treads} treads + {risers} risers, {sg['fw'] * 12:.1f}\" wide (count, not SF)"])
    # base and wall tile
    F.blank()
    F.section("BASE AND WALL TILE")
    F.header(["LEVEL", "ROOM", "ITEM", "LENGTH / AREA", "WASTE %", "ORDER QTY", "NOTES"])
    base = defaultdict(float)
    tile = 0.0
    for level in ("L1", "L2"):
        for r in M.rooms(level):
            if r.base in ("-", "") or r.num.startswith("ST"):
                continue
            g = M.clear_room_poly(r)
            perim = g.exterior.length
            dw = sum(o.w for o in M.doors(level) if _door_serves(o, r))
            dw += sum(o.w for o in M.OPENINGS if o.kind == "elevator" and o.level == level and r.num in ("100", "200"))
            if r.num in ("100", "200"):
                # corridor ends: storefront / link opening
                dw += 10 + 8 * IN if level == "L1" else 0
                dw += (12 - 2 * PL.CMUh) if level == "L1" else 0
            if r.num == "100A":
                dw += 2 * 32.0 + (12 - 2 * PL.CMUh)
            L = max(perim - dw, 0)
            base[r.base] += L
            F.add([level, f"{r.num} {r.name.replace(chr(10), ' ')}", r.base, round(L, 1), 0.05, "=D{r}*(1+E{r})",
                   "LF"], ff)
            if "CT-1" in r.walls:
                a = L * 7.0
                tile += a
                F.add([level, f"{r.num} {r.name.replace(chr(10), ' ')}", "CT-1 WAINSCOT", round(a, 1), 0.10,
                       "=D{r}*(1+E{r})", "SF (perimeter less doors x 7'-0\")"], ff)
    for k, v in sorted(base.items()):
        F.add(["", "TOTAL", k, round(v, 1), 0.05, "=D{r}*(1+E{r})", "LF"], ff, bold=True)
    F.add(["", "TOTAL", "CT-1", round(tile, 1), 0.10, "=D{r}*(1+E{r})", "SF"], ff, bold=True)

    # ===================================== DOORS =====================================
    D = Sheet(wb, "Doors-Frames-HW", ["DOOR", "LEVEL", "ROOM", "SIZE", "DOOR TYPE", "DOOR MATL", "FRAME TYPE",
                                      "FRAME MATL", "RATING", "HW SET", "REMARKS"],
              [9, 7, 9, 16, 9, 9, 9, 9, 9, 9, 40])
    for o in M.doors():
        size = ("PR " if o.pair else "") + f"{'3' if o.leaf == 3 else o.leaf}'-0\" x 7'-0\""
        D.add([o.id, o.level, o.room, size, o.type, o.dmat, o.frame, o.fmat, o.rating or "-", o.hw, o.remarks])
    D.blank()
    D.section("SUMMARY")
    def summarize(label, key):
        c = Counter(key(o) for o in M.doors())
        for k, n in sorted(c.items()):
            D.add([label, "", "", str(k), n])
    summarize("DOOR TYPE", lambda o: o.type)
    summarize("FRAME TYPE", lambda o: o.frame)
    summarize("HW SET", lambda o: o.hw)
    summarize("RATING", lambda o: o.rating or "NON-RATED")
    leaves = sum(2 if o.pair else 1 for o in M.doors())
    D.add(["TOTAL LEAVES", "", "", "", leaves])
    D.add(["TOTAL OPENINGS", "", "", "", len(M.doors())])

    # ===================================== WINDOWS =====================================
    W = Sheet(wb, "Windows-Storefront", ["TYPE", "DESCRIPTION", "MO W (FT)", "MO H (FT)", "COUNT", "AREA EACH (SF)",
                                         "TOTAL SF", "PERIMETER EACH (LF)", "TOTAL PERIM (LF)"],
              [8, 60, 10, 10, 8, 14, 12, 16, 16])
    wf = [None, None, "0.000", "0.000", "0", "#,##0.00", "#,##0.0", "#,##0.00", "#,##0.0"]
    c = Counter(o.type for o in M.OPENINGS if o.kind in ("window", "storefront", "louver"))
    for t in sorted(c):
        w, h = M.WINDOW_TYPES[t]["mo"]
        W.add([t, M.WINDOW_TYPES[t]["desc"], w, h, c[t], "=C{r}*D{r}", "=F{r}*E{r}", "=2*(C{r}+D{r})",
               "=H{r}*E{r}"], wf)
    W.blank()
    W.section("BY ELEVATION")
    W.header(["TYPE", "ELEVATION / LEVEL", "", "", "COUNT", "", "", "", ""])
    names = {"N": "NORTH", "S": "SOUTH", "E": "EAST", "W1": "WEST", "W2": "WEST", "W0": "WEST", "LN": "LINK NORTH",
             "LS": "LINK SOUTH"}
    cc = Counter((o.type, names[o.wall], o.level) for o in M.OPENINGS if o.kind in ("window", "storefront", "louver"))
    for (t, e, l), n in sorted(cc.items()):
        W.add([t, f"{e} / {l}", "", "", n])

    # ===================================== CARPENTRY =====================================
    C = Sheet(wb, "Carpentry", ["ITEM", "DESCRIPTION", "QTY", "UNIT", "NOTES"], [34, 70, 12, 8, 50])
    cf = [None, None, "#,##0.0", None, None]
    C.section("ROUGH CARPENTRY (TREATED / FRT BLOCKING)")
    par = 2 * (150 + 2 * OUT_FACE) + 2 * (72 + 2 * OUT_FACE)
    link_par = 2 * 36.0
    C.add(["PARAPET BLOCKING - MAIN ROOF", "PARAPET PERIMETER (LF), MULTIPLY BY NUMBER OF PLIES PER 8/A-501", par, "LF",
           "outside face perimeter"], cf)
    C.add(["PARAPET BLOCKING - LINK ROOF", "NORTH + SOUTH LINK PARAPETS", link_par, "LF", ""], cf)
    C.add(["ROOF-TO-WALL NAILER AT LINK", "MAIN BUILDING WEST WALL ABOVE LINK ROOF + EXIST. WALL", 2 * 12.0, "LF", ""], cf)
    perim = sum(2 * (o.w + (o.head - o.sill)) for o in M.OPENINGS
                if o.kind in ("window", "storefront") and o.wall in M.EXT_SEGS)
    C.add(["WINDOW / STOREFRONT PERIMETER BLOCKING", "SUM OF MO PERIMETERS (ALL W / SF TYPES)", perim, "LF",
           "if detailed on A-501"], cf)
    C.add(["RTU CURB NAILERS", "3 CURBS x 2 x (14' + 7')", 3 * 2 * 21.0, "LF", ""], cf)
    C.add(["ROOF HATCH CURB NAILER", "2 x (2.5' + 3.0')", 11.0, "LF", ""], cf)
    C.add(["FRT PLYWOOD BACKBOARDS", "3/4\" x 4' x 8': IDF 212 (3) + MECH/ELEC 112 (4)", 7, "SHT", ""], cf)
    C.section("ARCHITECTURAL CASEWORK (06 41 00)")
    roomcount = Counter(r.name for r in M.ROOMS)
    for name, items in M.CASEWORK_TYPICAL.items():
        n = roomcount.get(name, 0)
        for code, q in items.items():
            unit = "EA" if code == "CW-4" else "LF"
            C.add([f"{name} x {n}", f"{code}: {M.CASEWORK_TYPES[code]}", q * n, unit, f"{q} {unit} per room"], cf)
    tot = defaultdict(float)
    for name, items in M.CASEWORK_TYPICAL.items():
        for code, q in items.items():
            tot[code] += q * roomcount.get(name, 0)
    C.section("CASEWORK TOTALS")
    for code in sorted(tot):
        C.add([code, M.CASEWORK_TYPES[code], tot[code], "EA" if code == "CW-4" else "LF", ""], cf)

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, "CedarPrairie_ES_Addition_Practice_Quantity_Answer_Key.xlsx")
    wb.save(path)
    print("wrote", path)
    return path


def _door_serves(o, r):
    """does door o open into room r (for base deductions)"""
    if o.level != r.level:
        return False
    if o.room == r.num:
        return True
    corr = "100" if r.level == "L1" else "200"
    if r.num == corr and o.wall in ("B1", "B2", "C1", "C2", "C3", "C4"):
        return True
    if r.num == "100A" and o.id == "100A":
        return True
    if r.num in ("106", "206") and o.id in ("107", "207"):
        return True
    if r.num in ("110", "210") and o.id in ("110B", "210B"):
        return True
    return False


if __name__ == "__main__":
    build()
